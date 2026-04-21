"""
Generate MindCache Minor Project Report (DOCX)
Following IIIT Bhopal ECE Minor Project Report Format - 6th Semester
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
import os

doc = Document()

# ── Page margins ──
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(2.54)

# ── Helper Functions ──
def set_font(run, name='Times New Roman', size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading_text(text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold)
    return p

def add_body(text, size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8, first_indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(2)
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(first_indent)
    run = p.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    return p

def add_bullet(text, size=12, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    if bold_prefix:
        run = p.add_run(bold_prefix)
        set_font(run, size=size, bold=True)
        run = p.add_run(text)
        set_font(run, size=size)
    else:
        run = p.add_run(text)
        set_font(run, size=size)
    return p

def add_chapter_heading(chapter_num, title):
    doc.add_page_break()
    add_heading_text(f'CHAPTER {chapter_num}', size=18, bold=True, space_after=6)
    add_heading_text(title.upper(), size=16, bold=True, space_after=20)

def add_section_heading(text, size=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_font(run, size=size, bold=True)
    return p

def add_subsection_heading(text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_font(run, size=size, bold=True, italic=True)
    return p

# ═══════════════════════════════════════════════════════════════════
# TITLE PAGE
# ═══════════════════════════════════════════════════════════════════

for _ in range(3):
    doc.add_paragraph()

add_heading_text('MINDCACHE: AN AI-POWERED SELF-ORGANIZING', size=18, bold=True, space_after=4)
add_heading_text('KNOWLEDGE GRAPH SYSTEM FOR AUTONOMOUS', size=18, bold=True, space_after=4)
add_heading_text('MEMORY MANAGEMENT', size=18, bold=True, space_after=24)

add_heading_text('MINOR PROJECT REPORT', size=14, bold=True, space_after=12)
add_heading_text('Submitted in partial fulfillment for the award of the degree of', size=12, bold=False, space_after=8)
add_heading_text('BACHELOR OF TECHNOLOGY', size=14, bold=True, space_after=6)
add_heading_text('(Electronics and Communication Engineering)', size=12, bold=False, space_after=20)

add_heading_text('Submitted to', size=12, bold=False, space_after=6)
add_heading_text('INDIAN INSTITUTE OF INFORMATION TECHNOLOGY BHOPAL', size=13, bold=True, space_after=6)
add_heading_text('(Madhya Pradesh)', size=12, bold=False, space_after=20)

add_heading_text('Submitted by', size=12, bold=False, space_after=8)
add_heading_text('Mohd Faisal Hussain (Scholar Number)', size=12, bold=True, space_after=20)

add_heading_text('Under the supervision of', size=12, bold=False, space_after=8)
add_heading_text('Name of the Supervisor', size=12, bold=True, space_after=4)
add_heading_text('Assistant Professor (ECE)', size=12, bold=False, space_after=24)

add_heading_text('April 2026', size=12, bold=True, space_after=6)
add_heading_text('INDIAN INSTITUTE OF INFORMATION TECHNOLOGY BHOPAL', size=12, bold=True, space_after=6)

# ═══════════════════════════════════════════════════════════════════
# CERTIFICATE
# ═══════════════════════════════════════════════════════════════════
doc.add_page_break()
add_heading_text('CERTIFICATE', size=18, bold=True, space_after=24)

add_body(
    'This is to certify that the Minor Project Report entitled "MindCache: An AI-Powered '
    'Self-Organizing Knowledge Graph System for Autonomous Memory Management", submitted by '
    'Mohd Faisal Hussain in partial fulfillment of the requirements for the award of the degree '
    'of Bachelor of Technology in Electronics and Communication Engineering. This document is a '
    'comprehensive description of the work completed as part of the minor project evaluation.',
    space_after=20
)

add_body('Date:', space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)
add_body('', space_after=30)
add_heading_text('INDIAN INSTITUTE OF INFORMATION TECHNOLOGY BHOPAL', size=12, bold=True, space_after=6)

# ═══════════════════════════════════════════════════════════════════
# DECLARATION
# ═══════════════════════════════════════════════════════════════════
doc.add_page_break()
add_heading_text('DECLARATION', size=18, bold=True, space_after=24)

add_body(
    'I hereby declare that the following Minor Project Report entitled "MindCache: An AI-Powered '
    'Self-Organizing Knowledge Graph System for Autonomous Memory Management" presented in the '
    'partial fulfillment of the requirements for the award of the degree of Bachelor of Technology '
    'in Electronics and Communication Engineering. It is an authentic documentation of my original '
    'work carried out under the guidance of the project supervisor. The work has been carried out '
    'entirely at the Indian Institute of Information Technology, Bhopal. The project work presented '
    'has not been submitted in part or whole to award of any degree or professional diploma in any '
    'other institute or organization.',
    space_after=14
)

add_body(
    'I, with this, declare that the facts mentioned above are true to the best of my knowledge. '
    'In case of any unlikely discrepancy that may occur, I will be the one to take responsibility.',
    space_after=30
)

add_body('Mohd Faisal Hussain (Scholar Number)                                              Signature', space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)

# ═══════════════════════════════════════════════════════════════════
# AREA OF WORK
# ═══════════════════════════════════════════════════════════════════
doc.add_page_break()
add_heading_text('AREA OF WORK', size=18, bold=True, space_after=20)

add_body(
    'This project focuses on Artificial Intelligence, Natural Language Processing, and '
    'Graph-Based Knowledge Representation. The core challenge addressed is the autonomous '
    'organization and retrieval of unstructured digital context — a problem that sits at the '
    'intersection of information retrieval, LLM orchestration, and dynamic graph algorithms.',
    space_after=10
)

add_body(
    'MindCache is a full-stack AI system that passively captures a user\'s digital context '
    '(browsing activity, code sessions, research notes, error logs) through a browser extension, '
    'processes it through a multi-stage NLP pipeline, and autonomously organizes the extracted '
    'knowledge into a hierarchical, self-restructuring topic tree. Unlike traditional RAG '
    '(Retrieval-Augmented Generation) systems that rely on flat vector databases, MindCache '
    'maintains a living knowledge graph that continuously grooms, merges, and reorganizes itself '
    'to minimize redundancy and maximize retrieval precision.',
    space_after=10
)

add_body(
    'The system employs multiple specialized AI agents — an Input Denoiser for data preprocessing, '
    'a Memory Extractor for semantic parsing, a Decision State Analyzer for temporal reasoning, '
    'and a Tree Reorganizer for graph maintenance — each operating as independent components '
    'within a unified FastAPI-based backend architecture. The retrieval pipeline implements a '
    'novel 5-phase adaptive search strategy that combines vector embeddings with a custom BM25 '
    'keyword scorer, achieving 92% accuracy (47/51) on complex long-term memory evaluation benchmarks.',
    space_after=10
)

# ═══════════════════════════════════════════════════════════════════
# TABLE OF CONTENTS
# ═══════════════════════════════════════════════════════════════════
doc.add_page_break()
add_heading_text('TABLE OF CONTENTS', size=18, bold=True, space_after=20)

toc_items = [
    ('', 'Certificate', 'ii'),
    ('', 'Declaration', 'iii'),
    ('', 'Area of Work', 'iv'),
    ('', 'Table of Contents', 'v'),
    ('', 'List of Figures', 'vi'),
    ('', 'List of Tables', 'vii'),
    ('1', 'INTRODUCTION', ''),
    ('1.1', 'Background and Motivation', ''),
    ('1.2', 'Problem Statement', ''),
    ('1.3', 'Objectives', ''),
    ('1.4', 'Scope of the Project', ''),
    ('2', 'LITERATURE REVIEW', ''),
    ('2.1', 'Traditional Memory and Retrieval Systems', ''),
    ('2.2', 'LLM-Based Knowledge Extraction', ''),
    ('2.3', 'Graph-Based Knowledge Representation', ''),
    ('2.4', 'Hybrid Retrieval Strategies', ''),
    ('2.5', 'Gaps in Existing Work', ''),
    ('3', 'PROPOSED METHODOLOGY AND WORK DESCRIPTION', ''),
    ('3.1', 'System Overview', ''),
    ('3.2', 'Data Ingestion Layer', ''),
    ('3.3', 'Input Denoising Pipeline', ''),
    ('3.4', 'LLM-Based Memory Extraction', ''),
    ('3.5', 'Hierarchical Tree Organization', ''),
    ('3.6', 'Decision State Analysis', ''),
    ('3.7', 'Multi-Phase Retrieval Pipeline', ''),
    ('3.8', 'Fine-Tuning Data Curation', ''),
    ('4', 'PROPOSED ALGORITHMS', ''),
    ('4.1', 'Tri-State Input Classification Algorithm', ''),
    ('4.2', 'Code Skeletonizer State Machine', ''),
    ('4.3', 'Fuzzy-Hash Log Deduplication', ''),
    ('4.4', 'Tree Reorganization Multi-Phase Algorithm', ''),
    ('4.5', 'Hybrid BM25 + Vector Scoring Algorithm', ''),
    ('4.6', 'Adaptive Root Descent Algorithm', ''),
    ('5', 'PROPOSED FLOWCHART / BLOCK DIAGRAM', ''),
    ('5.1', 'System Architecture Block Diagram', ''),
    ('5.2', 'Data Ingestion Flow', ''),
    ('5.3', 'Memory Extraction Pipeline Flow', ''),
    ('5.4', 'Tree Reorganization Flow', ''),
    ('5.5', 'Retrieval Pipeline Flow', ''),
    ('6', 'TOOLS AND TECHNOLOGY USED', ''),
    ('7', 'CONCLUSION AND FUTURE SCOPE', ''),
    ('7.1', 'Conclusion', ''),
    ('7.2', 'Results and Evaluation', ''),
    ('7.3', 'Future Scope', ''),
    ('', 'REFERENCES', ''),
]

for num, title, page in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    if num and not '.' in num:
        run = p.add_run(f'Chapter {num}    {title}')
        set_font(run, size=12, bold=True)
    elif num:
        run = p.add_run(f'    {num}    {title}')
        set_font(run, size=12, bold=False)
    else:
        run = p.add_run(f'    {title}')
        set_font(run, size=12, bold='REFERENCE' in title.upper())

# ═══════════════════════════════════════════════════════════════════
# LIST OF FIGURES
# ═══════════════════════════════════════════════════════════════════
doc.add_page_break()
add_heading_text('LIST OF FIGURES', size=18, bold=True, space_after=20)

figures = [
    ('Figure 1.1', 'High-Level System Architecture of MindCache'),
    ('Figure 3.1', 'Complete System Pipeline Block Diagram'),
    ('Figure 3.2', 'Browser Extension Data Flow'),
    ('Figure 3.3', 'Tri-State Input Denoiser Classification Flow'),
    ('Figure 3.4', 'Code Skeletonizer State Machine Transitions'),
    ('Figure 3.5', 'Memory Extraction Pipeline'),
    ('Figure 3.6', 'Hierarchical Topic Tree Structure (Example)'),
    ('Figure 3.7', 'Tree Reorganization Multi-Phase Pipeline'),
    ('Figure 3.8', 'Decision State Analyzer Workflow'),
    ('Figure 5.1', 'System Architecture Block Diagram'),
    ('Figure 5.2', 'Data Ingestion Flow Diagram'),
    ('Figure 5.3', 'Memory Extraction Pipeline Flow'),
    ('Figure 5.4', 'Tree Reorganization Flow Diagram'),
    ('Figure 5.5', '5-Phase Retrieval Pipeline Flow'),
]

for fig_num, fig_title in figures:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f'{fig_num}: {fig_title}')
    set_font(run, size=12)

# ═══════════════════════════════════════════════════════════════════
# LIST OF TABLES
# ═══════════════════════════════════════════════════════════════════
doc.add_page_break()
add_heading_text('LIST OF TABLES', size=18, bold=True, space_after=20)

tables_list = [
    ('Table 3.1', 'Memory Types and Their Attributes'),
    ('Table 3.2', 'Tree Reorganization Phase Summary'),
    ('Table 3.3', 'Retrieval Pipeline Phase Descriptions'),
    ('Table 4.1', 'BM25 Scoring Parameters'),
    ('Table 6.1', 'Tools and Technologies Used'),
    ('Table 7.1', 'LongMemEval Benchmark Results'),
    ('Table 7.2', 'Comparison with Existing Systems'),
]

for tab_num, tab_title in tables_list:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(f'{tab_num}: {tab_title}')
    set_font(run, size=12)

# ═══════════════════════════════════════════════════════════════════
# CHAPTER 1: INTRODUCTION
# ═══════════════════════════════════════════════════════════════════
add_chapter_heading(1, 'Introduction')

add_section_heading('1.1 Background and Motivation')
add_body(
    'The rapid proliferation of digital information has created an unprecedented challenge for '
    'knowledge workers, researchers, and developers: the inability to efficiently recall and '
    'retrieve contextual information from their own digital activities. A software developer '
    'may spend hours debugging an issue they have already solved months ago, or a researcher '
    'may fail to connect insights from disparate reading sessions. The human brain is optimized '
    'for pattern recognition and creative synthesis, not for the faithful storage and retrieval '
    'of large volumes of factual, episodic, and procedural knowledge.',
    space_after=10, first_indent=1.27
)

add_body(
    'Recent advances in Large Language Models (LLMs) have made it possible to extract structured '
    'semantic information from unstructured natural language text with remarkable accuracy. '
    'Simultaneously, vector embedding techniques have enabled semantic similarity search at scale. '
    'However, most existing AI-powered knowledge management systems — commonly referred to as '
    '"second brain" applications — suffer from a fundamental architectural limitation: they treat '
    'knowledge storage as a flat, append-only vector database problem. Raw text chunks are '
    'embedded and stored in a vector database (such as ChromaDB, Pinecone, or Weaviate), and '
    'retrieval is performed via brute-force semantic similarity search.',
    space_after=10, first_indent=1.27
)

add_body(
    'This flat-storage paradigm fails to capture the hierarchical, temporal, and relational '
    'nature of real-world knowledge. Decisions evolve over time — what a user decided last month '
    'may have been superseded by new information. Topics have natural hierarchical relationships '
    '— "Machine Learning" subsumes "Neural Networks", which subsumes "Convolutional Neural Networks". '
    'These structural relationships are lost in a flat vector store.',
    space_after=10, first_indent=1.27
)

add_body(
    'MindCache was conceived to address this gap. Rather than treating knowledge as an unstructured '
    'bag of text chunks, MindCache autonomously organizes extracted memories into a self-restructuring '
    'hierarchical topic tree — a living knowledge graph that continuously grooms, merges, splits, '
    'and reorganizes itself to maintain semantic coherence as new information flows in.',
    space_after=10, first_indent=1.27
)

add_section_heading('1.2 Problem Statement')
add_body(
    'Existing AI-powered personal knowledge management systems rely on flat vector databases '
    'for storage and brute-force semantic similarity search for retrieval. This approach suffers '
    'from three critical limitations: (1) loss of hierarchical structure — the natural taxonomy '
    'of knowledge is destroyed when converted to flat embeddings; (2) absence of temporal reasoning '
    '— the system cannot determine whether a stored decision is still valid or has been superseded; '
    'and (3) poor scalability of retrieval precision — as the knowledge base grows, naive vector '
    'search returns increasingly noisy results because it lacks structural context to narrow the '
    'search space. This project proposes a solution that addresses all three limitations through '
    'a self-organizing hierarchical knowledge graph with autonomous graph maintenance and a '
    'multi-phase adaptive retrieval pipeline.',
    space_after=10, first_indent=1.27
)

add_section_heading('1.3 Objectives')
add_body('The primary objectives of this project are:', space_after=6)
add_bullet('To design and implement an intelligent preprocessing pipeline (Input Denoiser) that classifies and compresses raw input data (code, logs, natural language) to minimize token usage before LLM processing.')
add_bullet('To develop an LLM-based memory extraction engine capable of parsing unstructured text into typed, schema-constrained memory chains (episodic, knowledge, decision, preference).')
add_bullet('To engineer a self-organizing hierarchical topic tree with autonomous graph maintenance capabilities including root orthogonality enforcement, cycle resolution, sibling deduplication, fan-out limiting, and overloaded leaf auto-splitting.')
add_bullet('To implement a temporal Decision State Analyzer that autonomously tracks the validity of decision-type memories by cross-referencing them against surrounding episodic and knowledge context.')
add_bullet('To build a multi-phase adaptive retrieval pipeline combining vector embeddings with a custom BM25 keyword scorer for precise, context-aware information retrieval.')
add_bullet('To create a fine-tuning data curation pipeline that transforms real-world system usage into supervised training datasets for model specialization.')
add_bullet('To achieve at least 90% retrieval accuracy on complex long-term memory evaluation benchmarks.')

add_section_heading('1.4 Scope of the Project')
add_body(
    'The scope of MindCache encompasses the complete end-to-end pipeline from data ingestion '
    'to knowledge retrieval. The system includes: (a) a Chrome browser extension for passive '
    'data capture from browsing sessions; (b) a FastAPI backend server for request routing and '
    'queue management; (c) a multi-stage NLP preprocessing pipeline; (d) an LLM-powered memory '
    'extraction engine; (e) a 1500+ line graph restructuring engine for autonomous tree maintenance; '
    '(f) a Decision State Analyzer for temporal reasoning; (g) a 5-phase adaptive retrieval '
    'pipeline; (h) a fine-tuning data curation pipeline; and (i) a web-based visualization '
    'interface for graph exploration and query testing.',
    space_after=10, first_indent=1.27
)

# ═══════════════════════════════════════════════════════════════════
# CHAPTER 2: LITERATURE REVIEW
# ═══════════════════════════════════════════════════════════════════
add_chapter_heading(2, 'Literature Review')

add_section_heading('2.1 Traditional Memory and Retrieval Systems')
add_body(
    'The concept of augmenting human memory with digital systems dates back to Vannevar Bush\'s '
    'seminal 1945 essay "As We May Think", which proposed the Memex — a hypothetical device for '
    'storing and retrieving an individual\'s books, records, and communications. Modern implementations '
    'of this vision include personal knowledge management (PKM) tools such as Obsidian, Notion, '
    'and Roam Research, which organize information using manual tagging, linking, and hierarchical '
    'folder structures. While effective for users who invest time in explicit organization, these '
    'systems require significant manual effort and do not autonomously extract or structure knowledge '
    'from raw digital activity.',
    space_after=10, first_indent=1.27
)

add_section_heading('2.2 LLM-Based Knowledge Extraction')
add_body(
    'The emergence of Large Language Models, particularly the GPT family (OpenAI, 2020-2024), '
    'LLaMA (Meta, 2023), and Qwen (Alibaba, 2024), has revolutionized the ability to extract '
    'structured information from unstructured text. Techniques such as few-shot prompting, '
    'chain-of-thought reasoning, and constrained JSON output generation have enabled LLMs to '
    'function as reliable information extraction engines. Notable work in this area includes '
    'structured extraction frameworks like Instructor (Liu, 2023) and DSPy (Khattab et al., 2023), '
    'which provide programmatic interfaces for LLM-based data extraction with schema validation. '
    'MindCache builds upon these foundations by implementing a multi-pass extraction pipeline '
    'that uses Pydantic schema validation to ensure type-safe, structured output from local LLMs '
    'running via Ollama.',
    space_after=10, first_indent=1.27
)

add_section_heading('2.3 Graph-Based Knowledge Representation')
add_body(
    'Knowledge graphs have been extensively studied as a means of representing structured relationships '
    'between entities. Notable systems include Google\'s Knowledge Graph (Singhal, 2012), Wikidata '
    '(Vrandečić & Krötzsch, 2014), and domain-specific ontologies using RDF and OWL frameworks. '
    'In the personal knowledge management space, systems like MemPalace (2025) have adopted '
    'spatial metaphors (the ancient Greek method of loci) to organize memories into wings, rooms, '
    'and drawers. While MemPalace achieves impressive retrieval benchmarks (96.6% on LongMemEval), '
    'its graph structure is static — pre-defined spatial partitions do not adapt as the knowledge '
    'base evolves. MindCache addresses this limitation by implementing a dynamic, self-restructuring '
    'hierarchical topic tree that autonomously maintains its own organization.',
    space_after=10, first_indent=1.27
)

add_section_heading('2.4 Hybrid Retrieval Strategies')
add_body(
    'Modern information retrieval research has demonstrated that combining dense vector retrieval '
    '(semantic similarity via embeddings) with sparse retrieval (keyword matching via BM25) '
    'yields superior results compared to either approach alone. The BM25 algorithm (Robertson & '
    'Zaragoza, 2009) remains the gold standard for keyword-based retrieval, while dense retrieval '
    'using transformer-based embeddings (Karpukhin et al., 2020) captures semantic meaning beyond '
    'lexical overlap. Hybrid approaches such as those used in ColBERT (Khattab & Zaharia, 2020) '
    'and RRF (Reciprocal Rank Fusion) have shown significant improvements in retrieval precision. '
    'MindCache implements its own hybrid scoring mechanism with a configurable 60/40 weighting '
    'between vector similarity and a custom-built BM25 scorer.',
    space_after=10, first_indent=1.27
)

add_section_heading('2.5 Gaps in Existing Work')
add_body(
    'Despite significant progress in each of these individual areas, no existing system combines '
    'all four capabilities: (1) intelligent preprocessing that reduces token waste before LLM '
    'processing; (2) autonomous graph restructuring that maintains semantic coherence without '
    'human intervention; (3) temporal decision state tracking that distinguishes between active '
    'and superseded knowledge; and (4) multi-phase adaptive retrieval that leverages hierarchical '
    'graph structure to narrow the search space. MindCache is designed to fill this gap by '
    'integrating all four capabilities into a unified, end-to-end system.',
    space_after=10, first_indent=1.27
)

# ═══════════════════════════════════════════════════════════════════
# CHAPTER 3: PROPOSED METHODOLOGY AND WORK DESCRIPTION
# ═══════════════════════════════════════════════════════════════════
add_chapter_heading(3, 'Proposed Methodology and Work Description')

add_section_heading('3.1 System Overview')
add_body(
    'MindCache operates as a multi-layered AI pipeline with six core subsystems, each responsible '
    'for a distinct phase of the knowledge management lifecycle. The system accepts raw, unstructured '
    'digital context as input and produces a queryable, hierarchically organized knowledge graph '
    'as output. The six subsystems are: (1) Data Ingestion Layer, (2) Input Denoising Pipeline, '
    '(3) LLM-Based Memory Extraction, (4) Hierarchical Tree Organization and Maintenance, '
    '(5) Decision State Analysis, and (6) Multi-Phase Adaptive Retrieval.',
    space_after=10, first_indent=1.27
)

add_section_heading('3.2 Data Ingestion Layer')
add_body(
    'The data ingestion layer consists of a Chrome browser extension that passively captures '
    'the user\'s browsing context. The extension operates in two modes: (a) passive scraping mode, '
    'which monitors active browser tabs and extracts visible text content at configurable intervals; '
    'and (b) manual selection mode, which allows the user to highlight specific text passages for '
    'ingestion. Captured data is transmitted to the FastAPI backend server via REST API endpoints. '
    'The backend server (api_server.py, ~17KB) manages request queuing, authentication, and routes '
    'incoming data to the appropriate processing pipeline based on content type and priority.',
    space_after=10, first_indent=1.27
)

add_section_heading('3.3 Input Denoising Pipeline')
add_body(
    'Raw ingested data is typically noisy — it contains mixed code snippets, error log traces, '
    'HTML artifacts, and natural language text interspersed with formatting noise. Feeding this '
    'raw data directly to an LLM wastes valuable context window tokens and degrades extraction quality.',
    space_after=10, first_indent=1.27
)

add_subsection_heading('3.3.1 Tri-State Classification Router')
add_body(
    'The Input Denoiser implements a tri-state classification system that processes each line '
    'of input and classifies it into one of three categories: CODE, LOG, or TEXT. Classification '
    'is performed using two sets of compiled regular expression patterns — code_signals (detecting '
    'import statements, function definitions, class declarations, decorators, and bracket patterns) '
    'and log_signals (detecting timestamps, log levels, stack traces, and error codes). Lines '
    'matching code patterns are routed to the Code Compressor; lines matching log patterns are '
    'routed to the Log Compressor; remaining lines are treated as natural language text and passed '
    'through with minimal processing.',
    space_after=10, first_indent=1.27
)

add_subsection_heading('3.3.2 Code Skeletonizer (State Machine)')
add_body(
    'Code blocks are processed by a state-machine-based "skeletonizer" that strips function '
    'and method bodies while preserving architecturally significant elements: import statements, '
    'class declarations, function signatures (including decorators and type hints), module-level '
    'constants, and structural comments (TODO, FIXME, HACK annotations). The state machine '
    'tracks indentation depth to determine function body boundaries, emitting a "..." placeholder '
    'for stripped bodies. This preserves the architectural intent of the code without wasting '
    'tokens on implementation logic that is irrelevant to the semantic meaning of the user\'s activity.',
    space_after=10, first_indent=1.27
)

add_subsection_heading('3.3.3 Fuzzy-Hash Log Deduplicator')
add_body(
    'Error log segments are processed by a fuzzy-hash deduplication engine. The engine first '
    'stitches multi-line log entries (such as Python tracebacks) into complete units. It then '
    'creates a "fuzzy hash" of each log entry by masking dynamic values — UUIDs, IP addresses, '
    'hexadecimal memory addresses, timestamps, and numeric identifiers — with placeholder tokens. '
    'This normalized representation allows the deduplicator to identify structurally identical '
    'log entries that differ only in their dynamic values. Duplicate structures are collapsed '
    'into a single representative instance with a repetition count annotation. This compression '
    'technique typically achieves 40-60% token reduction before LLM processing.',
    space_after=10, first_indent=1.27
)

add_section_heading('3.4 LLM-Based Memory Extraction')
add_body(
    'Denoised input is processed by the Memory Extractor module, which uses LLMs (running locally '
    'via Ollama, primarily Qwen 1.7B/4B models) to parse unstructured text into typed memory chains. '
    'The extraction engine produces four types of memories:',
    space_after=6, first_indent=1.27
)

# Memory Types Table
table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
table.alignment = WD_TABLE_ALIGNMENT.CENTER

headers = ['Memory Type', 'Description', 'Example']
data = [
    ['Episodic', 'Events, sessions, debugging episodes with temporal context', '"Debugged the auth token refresh issue for 2 hours on March 10"'],
    ['Knowledge', 'Factual information, technical concepts, learned insights', '"FastAPI uses Starlette under the hood for async routing"'],
    ['Decision', 'Choices made with rationale, subject to temporal validity', '"Chose SQLite over PostgreSQL for single-user simplicity"'],
    ['Preference', 'User habits, opinions, workflow preferences', '"Always use type hints in Python function signatures"'],
]

for j, h in enumerate(headers):
    cell = table.rows[0].cells[j]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, row_data in enumerate(data):
    for j, cell_text in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_text

add_body('', space_after=6)
add_body('Table 3.1: Memory Types and Their Attributes', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

add_body(
    'Each extracted memory is tagged with a topic chain (e.g., ["Backend", "Database", "Migrations"]), '
    'a timestamp, and a confidence score. The extraction process uses Pydantic schema validation '
    '(via the SafeAI module) to ensure that LLM outputs conform to the expected JSON structure. '
    'Failed validations trigger automatic retry with error feedback injection into the prompt.',
    space_after=10, first_indent=1.27
)

add_section_heading('3.5 Hierarchical Tree Organization')
add_body(
    'The Tree Reorganizer (reorganize_tree.py, 1530 lines) is the engineering core of MindCache. '
    'It maintains a multi-rooted hierarchical topic tree where each root represents a top-level '
    'knowledge domain (e.g., "Computer Science", "Health", "Law"), and leaf nodes store actual '
    'memory references. The reorganizer operates in multiple phases:',
    space_after=6, first_indent=1.27
)

# Reorganization Phases Table
table2 = doc.add_table(rows=7, cols=3)
table2.style = 'Table Grid'
table2.alignment = WD_TABLE_ALIGNMENT.CENTER

headers2 = ['Phase', 'Name', 'Description']
data2 = [
    ['Phase 0', 'Root Orthogonality', 'Ensures top-level root domains are mutually exclusive and non-overlapping'],
    ['Phase 1', 'Global Bootstrap', 'Full tree restructuring for small databases (below threshold)'],
    ['Phase 2', 'Targeted Vector Grooming', 'Hybrid BM25 + vector similarity matching of ungroomed chains to existing branches'],
    ['Phase 3', 'Fan-Out Limiting', 'Clusters children when a parent exceeds 10 child nodes using LLM-driven grouping'],
    ['Phase 4', 'Overloaded Leaf Splitting', 'Auto-splits leaf nodes with 25+ memories into meaningful sub-categories'],
    ['Phase 5', 'Post-Processing', 'Sibling deduplication, orphan cleanup, cycle detection and resolution'],
]

for j, h in enumerate(headers2):
    cell = table2.rows[0].cells[j]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, row_data in enumerate(data2):
    for j, cell_text in enumerate(row_data):
        table2.rows[i+1].cells[j].text = cell_text

add_body('', space_after=6)
add_body('Table 3.2: Tree Reorganization Phase Summary', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

add_body(
    'A critical engineering challenge in the reorganizer is token-aware bin-packing: the system '
    'groups tree branches into optimally-sized batches (using a min_groups algorithm) so that each '
    'LLM restructuring call utilizes maximum context without exceeding the model\'s token limit. '
    'The reorganizer also maintains an embedding cache, updating or invalidating cached topic '
    'embeddings whenever the tree structure changes.',
    space_after=10, first_indent=1.27
)

add_section_heading('3.6 Decision State Analysis')
add_body(
    'The Decision State Analyzer (decision_analyzer.py) is a specialized LLM agent that addresses '
    'the temporal validity problem in knowledge management. When a user makes a decision (e.g., '
    '"Use PostgreSQL for the backend database"), that decision may later be superseded by new '
    'information (e.g., "Actually, switched to SQLite for simplicity"). The Decision Analyzer '
    'periodically scans decision-type memories and cross-references each decision against '
    'surrounding episodic and knowledge memories within a configurable time window.',
    space_after=10, first_indent=1.27
)

add_body(
    'For each decision, the analyzer classifies its current state into one of five categories: '
    'ACTIVE (still valid), SUPERSEDED (replaced by a newer decision), REJECTED (explicitly abandoned), '
    'CONDITIONAL (valid only under specific circumstances), or INACTIVE (no longer relevant due '
    'to context changes). This temporal reasoning capability ensures that retrieval results '
    'reflect the current state of the user\'s knowledge, not historical artifacts.',
    space_after=10, first_indent=1.27
)

add_section_heading('3.7 Multi-Phase Retrieval Pipeline')
add_body(
    'Querying the knowledge graph is performed through a 5-phase adaptive retrieval pipeline '
    '(active_path.py), designed to leverage the hierarchical structure of the topic tree for '
    'precise, context-aware retrieval:',
    space_after=6, first_indent=1.27
)

# Retrieval Phases Table
table3 = doc.add_table(rows=6, cols=3)
table3.style = 'Table Grid'
table3.alignment = WD_TABLE_ALIGNMENT.CENTER

headers3 = ['Phase', 'Component', 'Function']
data3 = [
    ['Phase 1', 'Context Bridge', 'Constructs a composite query vector by analyzing the current prompt and recent conversation history, detecting topic drift using cosine similarity thresholds'],
    ['Phase 2', 'Root Search', 'Identifies the 1-6 most relevant root domains by comparing the query vector against root embeddings'],
    ['Phase 3', 'Root Descent', 'Recursively walks the topic tree from selected roots, using hybrid 60/40 vector + BM25 scoring with an Adaptive Node Selector that escalates to LLM-based selection when vector scores are ambiguous'],
    ['Phase 4', 'Agentic Refiner', 'An LLM agent that classifies query intent (RECALL, EXPLAIN, COMPARE, PLAN, OVERVIEW), selects the 1-3 most relevant candidate topics, and decides retrieval depth (summary vs. leaf-level data)'],
    ['Phase 5', 'Database Fetch', 'Retrieves actual memory data from SQLite, applying decision-state filters to exclude superseded or rejected memories'],
]

for j, h in enumerate(headers3):
    cell = table3.rows[0].cells[j]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, row_data in enumerate(data3):
    for j, cell_text in enumerate(row_data):
        table3.rows[i+1].cells[j].text = cell_text

add_body('', space_after=6)
add_body('Table 3.3: Retrieval Pipeline Phase Descriptions', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

add_section_heading('3.8 Fine-Tuning Data Curation')
add_body(
    'To address the limitations of base LLMs in producing consistently schema-compliant structured '
    'output, MindCache includes a fine-tuning data curation pipeline (prepare_finetuning_data.py). '
    'This pipeline transforms real-world MindCache usage data — actual ingestion inputs paired '
    'with their corresponding validated extraction outputs — into supervised fine-tuning examples. '
    'The curated dataset (16MB+, stored as finetuning_dataset.json) is used to specialize small, '
    'resource-efficient local models (Qwen 1.7B) for the specific task of schema-constrained '
    'memory extraction, significantly improving extraction reliability without requiring larger, '
    'more expensive models.',
    space_after=10, first_indent=1.27
)

# ═══════════════════════════════════════════════════════════════════
# CHAPTER 4: PROPOSED ALGORITHMS
# ═══════════════════════════════════════════════════════════════════
add_chapter_heading(4, 'Proposed Algorithms')

add_section_heading('4.1 Tri-State Input Classification Algorithm')
add_body('The classification algorithm operates as follows:', space_after=6)
add_body(
    'INPUT: A sequence of text lines L = {l₁, l₂, ..., lₙ}\n'
    'OUTPUT: Classified segments S = {(type, content)} where type ∈ {CODE, LOG, TEXT}\n\n'
    'Algorithm:\n'
    '1. Initialize current_type = TEXT, buffer = []\n'
    '2. For each line lᵢ in L:\n'
    '   a. Compute code_score = Σ(match(lᵢ, pattern) for pattern in code_signals)\n'
    '   b. Compute log_score = Σ(match(lᵢ, pattern) for pattern in log_signals)\n'
    '   c. If code_score > log_score and code_score > 0: classify as CODE\n'
    '   d. Else if log_score > 0: classify as LOG\n'
    '   e. Else: classify as TEXT\n'
    '   f. If type changes from current_type: flush buffer as segment, update current_type\n'
    '   g. Append lᵢ to buffer\n'
    '3. Flush remaining buffer\n'
    '4. Route CODE segments → Code Skeletonizer\n'
    '5. Route LOG segments → Fuzzy-Hash Deduplicator\n'
    '6. Pass TEXT segments through with minimal processing',
    space_after=10, size=11
)

add_section_heading('4.2 Code Skeletonizer State Machine')
add_body(
    'The code skeletonizer operates as a finite state machine with three states:\n\n'
    'States: {SCAN, INSIDE_FUNCTION, SKIP_BODY}\n'
    'Transitions:\n'
    '• SCAN → INSIDE_FUNCTION: When a function/method signature (def, class) is detected\n'
    '• INSIDE_FUNCTION → SKIP_BODY: After emitting the signature, enter body-skip mode\n'
    '• SKIP_BODY → SCAN: When indentation returns to or below the function\'s base level\n\n'
    'Preserved Elements (always emitted):\n'
    '• Import statements (import, from...import)\n'
    '• Class and function signatures (with decorators and type hints)\n'
    '• Module-level constants and assignments\n'
    '• Structural comments (# TODO, # FIXME, # HACK, # NOTE)\n'
    '• Empty lines between top-level definitions (for readability)\n\n'
    'Stripped Elements (replaced with "..."):\n'
    '• Function/method bodies\n'
    '• Loop bodies at non-module level\n'
    '• Conditional branches at non-module level',
    space_after=10, size=11
)

add_section_heading('4.3 Fuzzy-Hash Log Deduplication')
add_body(
    'Algorithm for log deduplication:\n\n'
    '1. STITCH: Join multi-line log entries (e.g., Python tracebacks) into single records\n'
    '   - Detect continuation lines by indentation pattern and absence of timestamp prefix\n'
    '2. NORMALIZE: For each stitched log entry:\n'
    '   a. Replace UUIDs with <UUID>\n'
    '   b. Replace IP addresses with <IP>\n'
    '   c. Replace hex addresses (0x...) with <HEX>\n'
    '   d. Replace numeric sequences (>3 digits) with <NUM>\n'
    '   e. Replace ISO timestamps with <TIME>\n'
    '3. HASH: Compute fuzzy hash = hash(normalized_string)\n'
    '4. DEDUP: Use a hash table to detect duplicate structures\n'
    '   - If fuzzy_hash already seen: increment counter, skip entry\n'
    '   - If new: add to output, store hash\n'
    '5. ANNOTATE: Append "(×N)" count to deduplicated entries where N > 1\n'
    '6. OUTPUT: Return compressed log text',
    space_after=10, size=11
)

add_section_heading('4.4 Tree Reorganization Multi-Phase Algorithm')
add_body(
    'The tree reorganization algorithm is the most complex component of MindCache (1530 lines). '
    'The core algorithm operates as follows:\n\n'
    'INPUT: Topic tree T, set of new ungroomed chains C, embedding cache E\n'
    'OUTPUT: Restructured tree T\' with all chains properly placed\n\n'
    'Phase 0 — Root Orthogonality:\n'
    '  For each pair of roots (rᵢ, rⱼ): compute similarity(rᵢ, rⱼ)\n'
    '  If similarity > threshold: merge rⱼ into rᵢ (LLM decides merge target)\n'
    '  Dissolve generic roots (e.g., "General", "Miscellaneous")\n\n'
    'Phase 2 — Targeted Vector Grooming (for large DBs):\n'
    '  1. For each ungroomed chain c ∈ C:\n'
    '     a. Compute embedding(c)\n'
    '     b. Find top-k similar existing branches using cosine similarity\n'
    '     c. Compute BM25 score against candidate branch names\n'
    '     d. Rank by hybrid score = 0.6 * sim + 0.4 * bm25\n'
    '  2. Group chains into LLM batches using token-aware bin-packing:\n'
    '     a. Estimate token count for each chain + its candidate branch context\n'
    '     b. Use min_groups greedy algorithm to pack into minimum batches\n'
    '        where each batch < max_tokens\n'
    '  3. For each batch: invoke LLM to produce restructured subtree\n'
    '  4. Apply structural mutations (create, move, merge, rename nodes)\n'
    '  5. Detect and resolve cycles in the resulting graph\n'
    '  6. Update embedding cache for affected nodes\n\n'
    'Post-Processing:\n'
    '  • group_overgrown_children(): If any node has >10 children, cluster them\n'
    '  • split_overloaded_leaves(): If any leaf has >25 memories, subdivide\n'
    '  • Sibling deduplication: merge nodes with >0.9 embedding similarity\n'
    '  • Orphan cleanup: re-attach or remove nodes with no valid parent',
    space_after=10, size=11
)

add_section_heading('4.5 Hybrid BM25 + Vector Scoring Algorithm')
add_body(
    'The BM25 scorer is implemented from scratch (root_descent.py) following the Okapi BM25 '
    'formulation:\n\n'
    'BM25(q, d) = Σ IDF(qᵢ) · [f(qᵢ, d) · (k₁ + 1)] / [f(qᵢ, d) + k₁ · (1 - b + b · |d|/avgdl)]\n\n'
    'Where:\n'
    '• q = query terms, d = document (topic name + description)\n'
    '• f(qᵢ, d) = term frequency of query term qᵢ in document d\n'
    '• |d| = document length, avgdl = average document length across corpus\n'
    '• k₁ = 1.5 (term frequency saturation), b = 0.75 (length normalization)\n'
    '• IDF(qᵢ) = log((N - n(qᵢ) + 0.5) / (n(qᵢ) + 0.5) + 1)\n\n'
    'The hybrid score combines BM25 with vector cosine similarity:\n'
    '  hybrid_score = α · cosine_sim(query_vec, topic_vec) + (1-α) · normalize(BM25(q, d))\n'
    '  Where α = 0.6 (configurable via RetrievalConfig)',
    space_after=10, size=11
)

add_section_heading('4.6 Adaptive Root Descent Algorithm')
add_body(
    'The recursive descent algorithm navigates the topic tree from root to leaf:\n\n'
    '1. INPUT: query_vector, current_node, depth=0\n'
    '2. Fetch children of current_node\n'
    '3. For each child cᵢ: compute hybrid_score(query_vector, cᵢ)\n'
    '4. Filter: keep only children with hybrid_score > descent_threshold (0.20)\n'
    '5. SELECTION STRATEGY:\n'
    '   a. If |filtered_children| <= 3: use vector-based selection (top-k by score)\n'
    '   b. If |filtered_children| > 3 and scores are ambiguous (std_dev < 0.05):\n'
    '      escalate to LLM-based Adaptive Node Selector\n'
    '6. For each selected child:\n'
    '   a. If child is leaf: add to candidate list\n'
    '   b. If child is branch: RECURSE(query_vector, child, depth+1)\n'
    '7. OUTPUT: Ranked list of up to top_k (20) candidate topics',
    space_after=10, size=11
)

# ═══════════════════════════════════════════════════════════════════
# CHAPTER 5: PROPOSED FLOWCHART / BLOCK DIAGRAM
# ═══════════════════════════════════════════════════════════════════
add_chapter_heading(5, 'Proposed Flowchart / Block Diagram')

add_section_heading('5.1 System Architecture Block Diagram')
add_body(
    'The overall system architecture follows a layered pipeline design. Data flows from left '
    'to right through the following layers:\n\n'
    '┌──────────────┐    ┌──────────────┐    ┌──────────────┐    ┌──────────────┐    ┌──────────────┐\n'
    '│   Browser     │    │   FastAPI     │    │   Input       │    │   Memory      │    │   Topic      │\n'
    '│   Extension   │───▶│   Backend     │───▶│   Denoiser    │───▶│   Extractor   │───▶│   Tree       │\n'
    '│   (Chrome)    │    │   (Router)    │    │   (Tri-State) │    │   (LLM)       │    │   (SQLite)   │\n'
    '└──────────────┘    └──────────────┘    └──────────────┘    └──────────────┘    └──────┬───────┘\n'
    '                                                                                       │\n'
    '                                                                               ┌───────▼───────┐\n'
    '                                                                               │  Reorganizer   │\n'
    '                                                                               │  (1530 lines)  │\n'
    '                                                                               └───────┬───────┘\n'
    '                                                                                       │\n'
    '┌──────────────┐    ┌──────────────┐    ┌──────────────┐    ┌──────────────┐    ┌───────▼───────┐\n'
    '│   Context     │    │   Final       │    │   Agentic     │    │   Root        │    │   Root       │\n'
    '│   Output      │◀──│   DB Fetch    │◀──│   Refiner     │◀──│   Descent     │◀──│   Search     │\n'
    '│              │    │              │    │   (LLM)       │    │   (Hybrid)    │    │   (Vector)   │\n'
    '└──────────────┘    └──────────────┘    └──────────────┘    └──────────────┘    └──────────────┘',
    size=9, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10
)
add_body('Figure 5.1: System Architecture Block Diagram', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

add_section_heading('5.2 Data Ingestion Flow')
add_body(
    '┌─────────────┐     ┌──────────────┐     ┌─────────────┐     ┌──────────────┐\n'
    '│  User browses │     │  Extension    │     │  REST API    │     │  Ingestion   │\n'
    '│  web / codes  │────▶│  captures     │────▶│  POST to     │────▶│  queue       │\n'
    '│              │     │  context       │     │  /ingest     │     │  (async)     │\n'
    '└─────────────┘     └──────────────┘     └─────────────┘     └──────┬───────┘\n'
    '                                                                      │\n'
    '                                                              ┌───────▼───────┐\n'
    '                                                              │  Denoiser →    │\n'
    '                                                              │  Extractor →   │\n'
    '                                                              │  Reorganizer   │\n'
    '                                                              └───────────────┘',
    size=9, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10
)
add_body('Figure 5.2: Data Ingestion Flow Diagram', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

add_section_heading('5.3 Memory Extraction Pipeline Flow')
add_body(
    '┌──────────┐    ┌────────────┐    ┌──────────────┐    ┌────────────┐    ┌──────────┐\n'
    '│  Raw      │    │  Tri-State  │    │  CODE:       │    │  LLM       │    │  Pydantic │\n'
    '│  Input    │───▶│  Classify   │───▶│  Skeletonize │───▶│  Extract   │───▶│  Validate │\n'
    '│          │    │            │    │  LOG: Dedup   │    │  Memories  │    │  Schema   │\n'
    '│          │    │            │    │  TEXT: Pass    │    │            │    │           │\n'
    '└──────────┘    └────────────┘    └──────────────┘    └────────────┘    └─────┬────┘\n'
    '                                                                              │\n'
    '                                                                      ┌───────▼──────┐\n'
    '                                                                      │  Store in DB  │\n'
    '                                                                      │  + Reorganize │\n'
    '                                                                      └──────────────┘',
    size=9, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10
)
add_body('Figure 5.3: Memory Extraction Pipeline Flow', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

add_section_heading('5.4 Tree Reorganization Flow')
add_body(
    '┌───────────┐    ┌──────────────┐    ┌──────────────┐    ┌──────────────┐\n'
    '│  New       │    │  Phase 0:     │    │  Phase 1/2:  │    │  Post-       │\n'
    '│  Chains    │───▶│  Root Ortho   │───▶│  Grooming    │───▶│  Processing  │\n'
    '│  (input)   │    │  Check        │    │  (Bootstrap  │    │  (Fan-out,   │\n'
    '│           │    │              │    │   or Vector)  │    │  Split, Dedup│\n'
    '└───────────┘    └──────────────┘    └──────────────┘    └──────┬───────┘\n'
    '                                                                 │\n'
    '                                       ┌─────────────────────────▼──────┐\n'
    '                                       │  Update Embedding Cache        │\n'
    '                                       │  Resolve Cycles                │\n'
    '                                       │  Clean Orphans                 │\n'
    '                                       └────────────────────────────────┘',
    size=9, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10
)
add_body('Figure 5.4: Tree Reorganization Flow Diagram', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

add_section_heading('5.5 5-Phase Retrieval Pipeline Flow')
add_body(
    '┌──────────┐    ┌────────────┐    ┌────────────┐    ┌────────────┐    ┌────────────┐    ┌──────────┐\n'
    '│  User     │    │  Phase 1:   │    │  Phase 2:   │    │  Phase 3:   │    │  Phase 4:   │    │  Phase 5: │\n'
    '│  Query    │───▶│  Context    │───▶│  Root       │───▶│  Root       │───▶│  Agentic    │───▶│  DB       │\n'
    '│          │    │  Bridge     │    │  Search     │    │  Descent    │    │  Refiner    │    │  Fetch    │\n'
    '└──────────┘    └────────────┘    └────────────┘    └────────────┘    └────────────┘    └──────────┘\n'
    '                  │ Build          │ Find relevant    │ Walk tree       │ Intent classify  │ Retrieve\n'
    '                  │ query vector   │ root domains     │ hybrid scoring  │ Select 1-3       │ memories\n'
    '                  │ Detect drift   │ (vector sim)     │ BM25 + Vector   │ Decide depth     │ Apply\n'
    '                  │               │                  │ Adaptive select │                  │ filters',
    size=9, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10
)
add_body('Figure 5.5: 5-Phase Retrieval Pipeline Flow', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

# ═══════════════════════════════════════════════════════════════════
# CHAPTER 6: TOOLS AND TECHNOLOGY USED
# ═══════════════════════════════════════════════════════════════════
add_chapter_heading(6, 'Tools and Technology Used')

table6 = doc.add_table(rows=15, cols=3)
table6.style = 'Table Grid'
table6.alignment = WD_TABLE_ALIGNMENT.CENTER

headers6 = ['Category', 'Tool / Technology', 'Purpose']
tools_data = [
    ['Programming Language', 'Python 3.10+', 'Core system implementation'],
    ['Web Framework', 'FastAPI', 'REST API backend server, request routing, async handling'],
    ['Database', 'SQLite (via SQLAlchemy)', 'Persistent storage for topic tree, memories, embeddings'],
    ['LLM Runtime', 'Ollama', 'Local inference server for running LLMs without cloud API costs'],
    ['LLM Models', 'Qwen 1.7B / 4B, Gemma', 'Memory extraction, tree restructuring, decision analysis, retrieval refinement'],
    ['Embeddings', 'all-MiniLM-L6-v2 (Sentence Transformers)', 'Vector embedding generation for semantic similarity search'],
    ['NLP', 'Regex, Custom BM25, Tokenizer', 'Input classification, keyword scoring, text processing'],
    ['Schema Validation', 'Pydantic v2', 'Type-safe JSON schema enforcement for LLM outputs'],
    ['Frontend', 'HTML, CSS, JavaScript (Vanilla)', 'Web UI for graph visualization, workbench, explorer'],
    ['Browser Extension', 'Chrome Extension API (Manifest V3)', 'Passive data capture from browsing sessions'],
    ['Graph Visualization', 'Canvas 2D API (Custom)', 'Interactive tree visualization with pan/zoom/expand'],
    ['Version Control', 'Git + GitHub', 'Source code management and version tracking'],
    ['Fine-Tuning', 'Custom data pipeline + JSON dataset', 'Transforming usage data into supervised training examples'],
    ['Testing', 'Custom evaluation framework', 'LongMemEval benchmark integration for retrieval accuracy testing'],
]

for j, h in enumerate(headers6):
    cell = table6.rows[0].cells[j]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, row_data in enumerate(tools_data):
    for j, cell_text in enumerate(row_data):
        table6.rows[i+1].cells[j].text = cell_text

add_body('', space_after=6)
add_body('Table 6.1: Tools and Technologies Used', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# ═══════════════════════════════════════════════════════════════════
# CHAPTER 7: CONCLUSION AND FUTURE SCOPE
# ═══════════════════════════════════════════════════════════════════
add_chapter_heading(7, 'Conclusion and Future Scope')

add_section_heading('7.1 Conclusion')
add_body(
    'MindCache demonstrates that the predominant approach to personal AI memory systems — flat '
    'vector database storage with brute-force semantic search — is fundamentally insufficient '
    'for managing the hierarchical, temporal, and relational nature of real-world knowledge. '
    'By implementing a self-organizing hierarchical topic tree with autonomous graph maintenance, '
    'MindCache achieves a qualitative leap in how AI systems can structure and retrieve personal knowledge.',
    space_after=10, first_indent=1.27
)

add_body(
    'The key engineering contributions of this project are: (1) a tri-state Input Denoiser that '
    'reduces LLM token usage by up to 60% through intelligent preprocessing; (2) a 1530-line '
    'tree reorganization engine that autonomously maintains graph integrity through multi-phase '
    'structural grooming; (3) a Decision State Analyzer that introduces temporal reasoning into '
    'personal knowledge management; and (4) a 5-phase adaptive retrieval pipeline that combines '
    'vector and keyword-based scoring for precise hierarchical search.',
    space_after=10, first_indent=1.27
)

add_section_heading('7.2 Results and Evaluation')
add_body(
    'The system was evaluated using a custom benchmark derived from the LongMemEval framework, '
    'which tests the ability of a memory system to accurately retrieve specific information from '
    'long-term conversational context. MindCache achieved the following results:',
    space_after=6, first_indent=1.27
)

# Results Table
table7 = doc.add_table(rows=4, cols=3)
table7.style = 'Table Grid'
table7.alignment = WD_TABLE_ALIGNMENT.CENTER

headers7 = ['Metric', 'Score', 'Notes']
results_data = [
    ['Retrieval Accuracy', '92.16% (47/51)', 'Complex long-term memory evaluation queries'],
    ['Token Reduction (Denoiser)', 'Up to 60%', 'Measured on mixed code/log/text inputs'],
    ['Knowledge Graph Nodes', '24 root domains, 750+ topics', 'Autonomously organized from raw input'],
]

for j, h in enumerate(headers7):
    cell = table7.rows[0].cells[j]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, row_data in enumerate(results_data):
    for j, cell_text in enumerate(row_data):
        table7.rows[i+1].cells[j].text = cell_text

add_body('', space_after=6)
add_body('Table 7.1: LongMemEval Benchmark Results', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# Comparison Table
add_body('Comparison with existing systems:', space_after=6, bold=True)
table8 = doc.add_table(rows=4, cols=5)
table8.style = 'Table Grid'
table8.alignment = WD_TABLE_ALIGNMENT.CENTER

headers8 = ['System', 'Graph Structure', 'Self-Organizing', 'Temporal Reasoning', 'Retrieval Strategy']
comp_data = [
    ['Standard RAG', 'Flat (Vector DB)', 'No', 'No', 'Brute-force similarity'],
    ['MemPalace', 'Static spatial (Wing/Room)', 'No', 'Partial (valid_from/to)', 'Closet-first + fallback'],
    ['MindCache', 'Dynamic hierarchical tree', 'Yes (autonomous)', 'Yes (Decision Analyzer)', '5-phase adaptive hybrid'],
]

for j, h in enumerate(headers8):
    cell = table8.rows[0].cells[j]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for i, row_data in enumerate(comp_data):
    for j, cell_text in enumerate(row_data):
        table8.rows[i+1].cells[j].text = cell_text

add_body('', space_after=6)
add_body('Table 7.2: Comparison with Existing Systems', size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

add_section_heading('7.3 Future Scope')
add_body(
    'Several directions for future development have been identified:',
    space_after=6, first_indent=1.27
)
add_bullet('Multi-User Support: ', bold_prefix='')
add_body('Extending MindCache to support collaborative knowledge graphs where multiple users contribute to and query a shared knowledge base, with access control and attribution tracking.', first_indent=1.27, space_after=6)

add_bullet('Real-Time Streaming Ingestion: ', bold_prefix='')
add_body('Replacing the current batch-based ingestion with a real-time streaming pipeline using WebSockets, enabling sub-second memory extraction from live conversations.', first_indent=1.27, space_after=6)

add_bullet('Cross-Modal Memory: ', bold_prefix='')
add_body('Extending the memory extraction pipeline to support non-text modalities including images (screenshots, diagrams), audio (meeting recordings), and structured data (spreadsheets, databases).', first_indent=1.27, space_after=6)

add_bullet('Federated Learning for Privacy: ', bold_prefix='')
add_body('Implementing federated fine-tuning where multiple MindCache instances can collaboratively improve the extraction model without sharing raw personal data.', first_indent=1.27, space_after=6)

add_bullet('Advanced Graph Neural Networks: ', bold_prefix='')
add_body('Replacing the current embedding-based similarity scoring with Graph Neural Networks (GNNs) that can learn structural patterns in the topic tree for more accurate retrieval and reorganization.', first_indent=1.27, space_after=6)

add_bullet('Open-Source Release and Docker Deployment: ', bold_prefix='')
add_body('Packaging MindCache as a one-command Docker deployment for public release on GitHub, including comprehensive documentation, setup guides, and contribution guidelines.', first_indent=1.27, space_after=10)

# ═══════════════════════════════════════════════════════════════════
# REFERENCES
# ═══════════════════════════════════════════════════════════════════
doc.add_page_break()
add_heading_text('REFERENCES', size=18, bold=True, space_after=20)

references = [
    '[1] Bush, V. (1945). "As We May Think." The Atlantic Monthly, 176(1), 101-108.',
    '[2] Robertson, S., & Zaragoza, H. (2009). "The Probabilistic Relevance Framework: BM25 and Beyond." Foundations and Trends in Information Retrieval, 3(4), 333-389.',
    '[3] Singhal, A. (2012). "Introducing the Knowledge Graph." Google Official Blog.',
    '[4] Vrandečić, D., & Krötzsch, M. (2014). "Wikidata: A Free Collaborative Knowledgebase." Communications of the ACM, 57(10), 78-85.',
    '[5] Vaswani, A., et al. (2017). "Attention Is All You Need." Advances in Neural Information Processing Systems (NeurIPS).',
    '[6] Devlin, J., et al. (2019). "BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding." NAACL-HLT.',
    '[7] Karpukhin, V., et al. (2020). "Dense Passage Retrieval for Open-Domain Question Answering." EMNLP.',
    '[8] Khattab, O., & Zaharia, M. (2020). "ColBERT: Efficient and Effective Passage Search via Contextualized Late Interaction over BERT." SIGIR.',
    '[9] Brown, T., et al. (2020). "Language Models are Few-Shot Learners." NeurIPS.',
    '[10] Touvron, H., et al. (2023). "LLaMA: Open and Efficient Foundation Language Models." arXiv preprint.',
    '[11] Khattab, O., et al. (2023). "DSPy: Compiling Declarative Language Model Calls into Self-Improving Pipelines." arXiv preprint.',
    '[12] Liu, J. (2023). "Instructor: Structured Extraction with LLMs." GitHub Repository.',
    '[13] Bai, J., et al. (2024). "Qwen Technical Report." arXiv preprint.',
    '[14] Reimers, N., & Gurevych, I. (2019). "Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks." EMNLP.',
    '[15] MemPalace Contributors. (2025). "MemPalace: A Spatial Memory System for AI Assistants." GitHub Repository.',
]

for ref in references:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.27)
    p.paragraph_format.first_line_indent = Cm(-1.27)
    run = p.add_run(ref)
    set_font(run, size=11)

# ═══════════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════════
output_path = r'e:\MindCache\MindCache_Minor_Project_Report.docx'
doc.save(output_path)
print(f"Report generated successfully: {output_path}")
print(f"File size: {os.path.getsize(output_path) / 1024:.1f} KB")
